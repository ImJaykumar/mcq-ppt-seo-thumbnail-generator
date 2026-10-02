import io
import os
import pandas as pd
from PIL import Image, ImageDraw, ImageFont
import streamlit as st
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

# -------------------------------------------------------------
# PAGE CONFIGURATION & STYLING
# -------------------------------------------------------------
st.set_page_config(
    page_title="MCQ PPT, SEO & Thumbnail Generator",
    page_icon="🎬",
    layout="wide"
)

# -------------------------------------------------------------
# THEMES COLOR PALETTES DEFINITION FOR PPT
# -------------------------------------------------------------
THEMES = {
    "Ultra Dark Navy": {
        "BG_COLOR": RGBColor(15, 23, 42),       # Deep Slate (#0F172A)
        "CARD_BG": RGBColor(30, 41, 59),        # Dark Card Fill (#1E293B)
        "ACCENT_BLUE": RGBColor(56, 189, 248),  # Neon Cyan (#38BDF8)
        "ACCENT_GREEN": RGBColor(74, 222, 128), # Neon Green (#4ADE80)
        "ACCENT_ORANGE": RGBColor(251, 146, 60),# Warm Amber (#FB923C)
        "TEXT_WHITE": RGBColor(248, 250, 252),  # Pure Light (#F8FAFC)
        "TEXT_MUTED": RGBColor(148, 163, 184),  # Slate Grey (#94A3B8)
        "BORDER_COLOR": RGBColor(51, 65, 85)   # Border (#334155)
    },
    "Clean Light Mode": {
        "BG_COLOR": RGBColor(241, 245, 249),    # Light Grey (#F1F5F9)
        "CARD_BG": RGBColor(255, 255, 255),     # Pure White (#FFFFFF)
        "ACCENT_BLUE": RGBColor(37, 99, 235),   # Royal Blue (#2563EB)
        "ACCENT_GREEN": RGBColor(22, 163, 74),  # Forest Green (#16A34A)
        "ACCENT_ORANGE": RGBColor(234, 88, 12), # Orange (#EA580C)
        "TEXT_WHITE": RGBColor(15, 23, 42),     # Deep Dark Text (#0F172A)
        "TEXT_MUTED": RGBColor(100, 116, 139),  # Muted Dark Grey (#64748B)
        "BORDER_COLOR": RGBColor(203, 213, 225) # Soft Border (#CBD5E1)
    },
    "Emerald Forest": {
        "BG_COLOR": RGBColor(6, 78, 59),       # Deep Forest (#064E3B)
        "CARD_BG": RGBColor(4, 120, 87),       # Emerald Card (#047857)
        "ACCENT_BLUE": RGBColor(110, 231, 183), # Mint Green (#6EE7B7)
        "ACCENT_GREEN": RGBColor(253, 224, 71), # Bright Yellow Accent (#FDE047)
        "ACCENT_ORANGE": RGBColor(249, 115, 22),# Amber (#F97316)
        "TEXT_WHITE": RGBColor(240, 253, 244), # Soft Mint White (#F0FDF4)
        "TEXT_MUTED": RGBColor(167, 243, 208), # Light Mint Grey (#A7F3D0)
        "BORDER_COLOR": RGBColor(16, 185, 129) # Emerald Border (#10B981)
    },
    "Minimalist Slate": {
        "BG_COLOR": RGBColor(24, 24, 27),       # Zinc Black (#18181B)
        "CARD_BG": RGBColor(39, 39, 42),       # Dark Slate Card (#27272A)
        "ACCENT_BLUE": RGBColor(161, 161, 170),# Neutral Muted White (#A1A1AA)
        "ACCENT_GREEN": RGBColor(34, 197, 94),  # Standard Green (#22C55E)
        "ACCENT_ORANGE": RGBColor(245, 158, 11),# Amber (#F59E0B)
        "TEXT_WHITE": RGBColor(250, 250, 250), # Pure White (#FAFAFA)
        "TEXT_MUTED": RGBColor(161, 161, 170), # Zinc Grey (#A1A1AA)
        "BORDER_COLOR": RGBColor(63, 63, 70)   # Border Slate (#3F3F46)
    }
}


# -------------------------------------------------------------
# SAFE YOUTUBE THUMBNAIL GENERATOR FUNCTION
# -------------------------------------------------------------
def generate_youtube_thumbnail(subject_text, total_q_text, tag_text="KATHAN & KUT PATTERN"):
    """Creates a high-contrast HD 1280x720 YouTube Thumbnail safely."""
    width, height = 1280, 720
    img = Image.new("RGB", (width, height), color="#0F172A")
    draw = ImageDraw.Draw(img)

    # Decorative Borders
    draw.rectangle([30, 30, width - 30, height - 30], outline="#38BDF8", width=5)
    draw.rectangle([50, 50, width - 50, height - 50], outline="#334155", width=2)

    # Header Badge Box
    draw.rounded_rectangle([80, 80, 550, 150], radius=15, fill="#F59E0B")
    
    # Safe Default Font Fallback (Prevents Crash/Blank Screen)
    font_header = font_subject = font_qcount = font_tag = ImageFont.load_default()
    try:
        font_header = ImageFont.truetype("arial.ttf", 36)
        font_subject = ImageFont.truetype("arialbd.ttf", 60)
        font_qcount = ImageFont.truetype("arialbd.ttf", 68)
        font_tag = ImageFont.truetype("arialbd.ttf", 38)
    except Exception:
        pass

    # Draw Text Elements
    draw.text((100, 95), "CGPSC & VYAPAM 2026", fill="#0F172A", font=font_header)
    draw.text((80, 190), subject_text.upper(), fill="#38BDF8", font=font_subject)

    # Question Count Banner
    draw.rounded_rectangle([80, 300, 850, 430], radius=20, fill="#1E293B", outline="#4ADE80", width=4)
    draw.text((110, 330), f"TOP {total_q_text} EXPECTED MCQs", fill="#4ADE80", font=font_qcount)

    # Strategy Tag
    draw.rounded_rectangle([80, 470, 900, 570], radius=15, fill="#EF4444")
    draw.text((110, 495), f"🔥 {tag_text.upper()}", fill="#FFFFFF", font=font_tag)

    # Right Accent Element
    draw.ellipse([980, 180, 1200, 400], fill="#1E293B", outline="#38BDF8", width=6)
    draw.text((1015, 260), "100%", fill="#FFFFFF", font=font_qcount)
    draw.text((1020, 330), "TARGET", fill="#F59E0B", font=font_header)

    img_io = io.BytesIO()
    img.save(img_io, format="PNG")
    img_io.seek(0)
    return img_io


# -------------------------------------------------------------
# PPT HELPER FUNCTIONS
# -------------------------------------------------------------
def add_text_to_shape(shape, text, font_size, color, bold=False, space_after=0):
    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0] if tf.paragraphs else tf.add_paragraph()
    p.text = str(text)
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    if space_after:
        p.space_after = Pt(space_after)
    return tf


def add_brand_and_timer(slide, stage, theme, logo_bytes=None):
    logo_left = Inches(11.3)
    logo_top = Inches(0.3)

    if logo_bytes is not None:
        logo_stream = io.BytesIO(logo_bytes)
        slide.shapes.add_picture(logo_stream, logo_left, logo_top, width=Inches(1.2))
    else:
        logo_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, logo_left, logo_top, Inches(1.2), Inches(0.4))
        logo_box.fill.solid()
        logo_box.fill.fore_color.rgb = theme["CARD_BG"]
        logo_box.line.color.rgb = theme["ACCENT_BLUE"]
        add_text_to_shape(logo_box, "CG EXAMS", 9, theme["TEXT_WHITE"], bold=True)

    if stage == 1:
        timer_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.6), logo_top, Inches(1.5), Inches(0.4))
        timer_box.fill.solid()
        timer_box.fill.fore_color.rgb = theme["CARD_BG"]
        timer_box.line.color.rgb = theme["ACCENT_ORANGE"]
        timer_box.line.width = Pt(1.5)
        add_text_to_shape(timer_box, "⏱️ TIMER: 10s", 10, theme["ACCENT_ORANGE"], bold=True)


def generate_seo_text(df):
    total_q = len(df)
    first_sub = str(df['Subject'].iloc[0]).strip() if 'Subject' in df.columns and len(df) > 0 else "CG General Knowledge"

    title_main = f"CGPSC & CG Vyapam 2026 | {first_sub} Top {total_q} MCQs (Kathan-Karan & Kut Pattern Special)"
    title_alt_1 = f"🔥 {first_sub} Top {total_q} Most Expected Questions | CGPSC Prelims 2026 & Vyapam Exams"

    timestamps_list = ["00:00 - Intro & Exam Pattern Strategy"]
    curr_sec = 20
    for i, r in df.iterrows():
        m, s = divmod(curr_sec, 60)
        q_type_str = str(r.get('QuestionType', 'MCQ')).strip()
        q_preview = str(r.get('Question', '')).strip()[:40]
        timestamps_list.append(f"{m:02d}:{s:02d} - Q{i + 1}: [{q_type_str}] {q_preview}...")
        curr_sec += 90

    timestamps_str = "\n".join(timestamps_list)

    keywords = [
        "CGPSC 2026", "CG Vyapam 2026", f"CGPSC {first_sub} MCQ", f"CG Vyapam {first_sub}",
        "Chhattisgarh General Knowledge", "CGPSC Kathan Karan Question", "CG Vyapam Model Paper",
        "CGPSC Prelims Test Series", "CG Hosteller Superintendent", "CG Sub Inspector Exam",
        "CGPSC Previous Year Questions", "Chhattisgarh GS Strategy"
    ]

    hashtags = [
        "#CGPSC2026", "#CGVyapam", f"#{first_sub.replace(' ', '')}",
        "#ChhattisgarhGK", "#CGPSC_MCQ", "#CG_Exams_2026"
    ]

    tags_str = ", ".join(keywords)
    hashtags_str = " ".join(hashtags)

    lines = [
        "====================================================================",
        "📌 YOUTUBE TITLE OPTIONS (Choose One)",
        "====================================================================",
        f"Primary Title : {title_main}",
        f"Alt Title (High CTR) : {title_alt_1}",
        "",
        "====================================================================",
        "📝 HIGH-CONVERTING YOUTUBE DESCRIPTION",
        "====================================================================",
        f"Welcome Aspirants! Is video me humne CGPSC aur CG Vyapam 2026 exams ke liye {first_sub} ke Sabse Important aur Expected {total_q} Questions cover kiye hain. Is me Kathan-Karan (Statement-Reason) aur Kut (Matching) pattern wale questions ko Elimination Trick ke sath solve kiya gaya hai.",
        "",
        "🎯 Useful For:",
        "- CGPSC State Service Prelims 2026",
        "- CG Vyapam (Hostel Warden, RI, Patwari, ADDO, SI)",
        "- Other Chhattisgarh State Competitive Exams",
        "",
        "👇 Fast Navigation / Video Timestamps:",
        timestamps_str,
        "",
        "--------------------------------------------------------------------",
        "📚 Important Resources & Playlists:",
        "👉 Join Telegram Channel for Free PDF: https://t.me/YourChannelLink",
        "👉 CGPSC Complete Strategy Playlist: https://youtube.com/playlist?list=YourPlaylistID",
        "",
        "====================================================================",
        "🏷️ YOUTUBE TAGS (Copy & Paste directly into YouTube Studio Tags Box)",
        "====================================================================",
        tags_str,
        "",
        "====================================================================",
        "#️⃣ HASHTAGS (Add at the end of Description)",
        "====================================================================",
        hashtags_str,
        "",
        "====================================================================",
        "💬 PINNED COMMENT (Copy & Paste in your Video Comments Section)",
        "====================================================================",
        f"Aapka Total {total_q} me se kitna score raha? Comments me zaroor batayein! 👇",
        "PDF Notes download karne ke liye hamare Telegram Group ko join karein: https://t.me/YourChannelLink",
        "Video pasand aaya ho toh LIKE aur SHARE zaroor karein! 🔔"
    ]

    return "\n".join(lines)


def build_presentation(df, theme, logo_bytes=None):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]
    total_q = len(df)

    for index, row in df.iterrows():
        subject = str(row.get('Subject', 'CG EXAMS')).strip().upper() or 'CG EXAMS'
        q_type = str(row.get('QuestionType', 'MCQ')).strip().upper() or 'MCQ'
        q_text = str(row.get('Question', '')).strip()
        correct_opt = str(row.get('CorrectAnswer', '')).strip()

        options = [
            ("A", str(row.get('OptionA', '')).strip()),
            ("B", str(row.get('OptionB', '')).strip()),
            ("C", str(row.get('OptionC', '')).strip()),
            ("D", str(row.get('OptionD', '')).strip())
        ]

        for stage in [1, 2]:
            slide = prs.slides.add_slide(blank_layout)

            bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
            bg.fill.solid()
            bg.fill.fore_color.rgb = theme["BG_COLOR"]
            bg.line.fill.background()

            header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(8.5), Inches(0.5))
            add_text_to_shape(
                header_box,
                f"CGPSC & CG VYAPAM SPECIAL  |  {subject} [{q_type}]  |  Q{index + 1}/{total_q}",
                13, theme["ACCENT_BLUE"], bold=True
            )

            add_brand_and_timer(slide, stage, theme, logo_bytes)

            q_box_w = Inches(7.4) if stage == 2 else Inches(11.733)
            q_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.0), q_box_w, Inches(2.3))
            q_card.fill.solid()
            q_card.fill.fore_color.rgb = theme["CARD_BG"]
            q_card.line.color.rgb = theme["BORDER_COLOR"]
            add_text_to_shape(q_card, q_text, 14, theme["TEXT_WHITE"], bold=True)

            opt_width = Inches(3.6) if stage == 2 else Inches(5.7)
            opt_positions = [
                (Inches(0.8), Inches(3.5)),
                (Inches(4.6) if stage == 2 else Inches(6.8), Inches(3.5)),
                (Inches(0.8), Inches(4.5)),
                (Inches(4.6) if stage == 2 else Inches(6.8), Inches(4.5))
            ]

            for idx, (label, opt_val) in enumerate(options):
                x, y = opt_positions[idx]
                opt_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, opt_width, Inches(0.85))
                opt_card.fill.solid()
                opt_card.fill.fore_color.rgb = theme["CARD_BG"]

                is_correct = stage == 2 and (label == correct_opt or opt_val == correct_opt or label in correct_opt)
                if is_correct:
                    opt_card.line.color.rgb = theme["ACCENT_GREEN"]
                    opt_card.line.width = Pt(2.5)
                else:
                    opt_card.line.color.rgb = theme["BORDER_COLOR"]
                    opt_card.line.width = Pt(1)

                add_text_to_shape(opt_card, f"({label})  {opt_val}", 13, theme["TEXT_WHITE"])

            if stage == 2:
                exp_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.4), Inches(1.0), Inches(4.133), Inches(5.5))
                exp_card.fill.solid()
                exp_card.fill.fore_color.rgb = theme["CARD_BG"]
                exp_card.line.color.rgb = theme["ACCENT_BLUE"]
                exp_card.line.width = Pt(1.5)

                tf_exp = exp_card.text_frame
                tf_exp.word_wrap = True

                p_ans = tf_exp.paragraphs[0]
                p_ans.text = f"✔ CORRECT: {correct_opt}"
                p_ans.font.size = Pt(14)
                p_ans.font.bold = True
                p_ans.font.color.rgb = theme["ACCENT_GREEN"]
                p_ans.space_after = Pt(10)

                elim_trick = str(row.get('EliminationTrick', '')).strip()
                if elim_trick:
                    p_tr_h = tf_exp.add_paragraph()
                    p_tr_h.text = "⚡ ELIMINATION TRICK"
                    p_tr_h.font.size = Pt(11)
                    p_tr_h.font.bold = True
                    p_tr_h.font.color.rgb = theme["ACCENT_BLUE"]

                    p_tr = tf_exp.add_paragraph()
                    p_tr.text = elim_trick
                    p_tr.font.size = Pt(12)
                    p_tr.font.color.rgb = theme["TEXT_MUTED"]
                    p_tr.space_after = Pt(10)

                key_points = str(row.get('KeyPoints', '')).strip()
                if key_points:
                    p_kp_h = tf_exp.add_paragraph()
                    p_kp_h.text = "📌 KEY REVISION POINTS"
                    p_kp_h.font.size = Pt(11)
                    p_kp_h.font.bold = True
                    p_kp_h.font.color.rgb = theme["ACCENT_BLUE"]

                    for pt in key_points.split('\n'):
                        if pt.strip():
                            p_pt = tf_exp.add_paragraph()
                            p_pt.text = pt.strip()
                            p_pt.font.size = Pt(11)
                            p_pt.font.color.rgb = theme["TEXT_WHITE"]

            max_width = 11.733
            progress_ratio = (index + (0.5 if stage == 1 else 1.0)) / total_q
            current_progress_width = max_width * progress_ratio

            t_track = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.75), Inches(max_width), Inches(0.08))
            t_track.fill.solid()
            t_track.fill.fore_color.rgb = theme["BORDER_COLOR"]
            t_track.line.fill.background()

            t_fill = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.75), Inches(current_progress_width), Inches(0.08))
            t_fill.fill.solid()
            t_fill.fill.fore_color.rgb = theme["ACCENT_BLUE"] if stage == 1 else theme["ACCENT_GREEN"]
            t_fill.line.fill.background()

    ppt_io = io.BytesIO()
    prs.save(ppt_io)
    ppt_io.seek(0)
    return ppt_io


# -------------------------------------------------------------
# STREAMLIT UI LAYOUT
# -------------------------------------------------------------
st.title("🎬 MCQ PPT, SEO & Thumbnail Automation Suite")
st.markdown("Upload your Excel file to generate **PPT Deck**, **YouTube SEO**, and an **Auto HD Thumbnail**.")

# Sidebar Controls
st.sidebar.header("🎨 Theme Customization")
selected_theme_name = st.sidebar.selectbox(
    "Choose PPT Color Theme:",
    options=list(THEMES.keys()),
    index=0
)
chosen_theme = THEMES[selected_theme_name]

st.sidebar.markdown("---")
st.sidebar.header("🖼️ Thumbnail Customization")
custom_badge_tag = st.sidebar.text_input("Thumbnail Badge Tag:", value="Kathan & Kut Pattern")

st.sidebar.markdown("---")
st.sidebar.header("📁 Upload Files")
uploaded_excel = st.sidebar.file_uploader("Upload Excel File (.xlsx)", type=["xlsx"])
uploaded_logo = st.sidebar.file_uploader("Upload Channel Logo (Optional PNG)", type=["png", "jpg", "jpeg"])

if uploaded_excel is not None:
    try:
        df = pd.read_excel(uploaded_excel).fillna("")
        df.columns = df.columns.astype(str).str.strip()

        st.success(f"✅ Excel file uploaded! Total Questions: **{len(df)}**")
        st.info(f"🎨 Active Theme Selected: **{selected_theme_name}**")

        # Preview Section
        with st.expander("🔍 Preview Uploaded Questions Data"):
            st.dataframe(df.head(10))

        logo_bytes = uploaded_logo.getvalue() if uploaded_logo is not None else None

        # Process Button
        if st.button("🚀 Generate All Assets (PPT, SEO & Thumbnail)", type="primary"):
            with st.spinner("Generating Assets... Please wait."):
                # 1. Generate PPTX
                ppt_io = build_presentation(df, theme=chosen_theme, logo_bytes=logo_bytes)
                
                # 2. Generate SEO Text
                seo_content = generate_seo_text(df)

                # 3. Generate YouTube Thumbnail
                first_sub = str(df['Subject'].iloc[0]).strip() if 'Subject' in df.columns and len(df) > 0 else "CG General Knowledge"
                thumb_io = generate_youtube_thumbnail(
                    subject_text=first_sub,
                    total_q_text=str(len(df)),
                    tag_text=custom_badge_tag
                )

                st.balloons()
                st.success("🎉 All assets generated successfully!")

                # Preview Thumbnail Image
                st.subheader("🖼️ Auto-Generated YouTube Thumbnail Preview")
                st.image(thumb_io, caption="HD 1280x720 Auto Thumbnail", use_container_width=True)

                st.markdown("---")
                st.subheader("📥 Download Generated Files")

                col1, col2, col3 = st.columns(3)

                with col1:
                    st.download_button(
                        label="📥 Download PPT Deck (.pptx)",
                        data=ppt_io,
                        file_name=f"CGPSC_MCQ_Deck_{selected_theme_name.replace(' ', '_')}.pptx",
                        mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
                        use_container_width=True
                    )

                with col2:
                    st.download_button(
                        label="🖼️ Download HD Thumbnail (.png)",
                        data=thumb_io,
                        file_name="YouTube_Thumbnail_HD.png",
                        mime="image/png",
                        use_container_width=True
                    )

                with col3:
                    st.download_button(
                        label="📄 Download SEO Output (.txt)",
                        data=seo_content,
                        file_name="youtube_seo_output_v2.txt",
                        mime="text/plain",
                        use_container_width=True
                    )

                # SEO Text Preview
                st.markdown("---")
                st.subheader("📋 YouTube SEO Text Preview")
                st.code(seo_content, language="text")

    except Exception as e:
        st.error(f"❌ Error processing file: {e}")
else:
    st.info("👈 Upload an Excel file from the sidebar to generate PPT, SEO, and Thumbnail.")